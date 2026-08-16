"""
export_service の単体テスト（HTTP経由ではなく関数を直接呼ぶ）。
"""
from __future__ import annotations

from datetime import date

import pytest
import pytest_asyncio

from app.models.department import Department, WorkRuleConfig
from app.models.employee import Employee
from app.models.leave_request import LeaveRequest
from app.models.schedule import ShiftAssignment, ShiftSchedule
from app.models.work_pattern import WorkPattern, WorkPatternGroup
from app.services.export_service import (
    export_schedule_csv_long,
    export_schedule_xlsx,
)


@pytest_asyncio.fixture
async def schedule_full(client: tuple) -> ShiftSchedule:
    """1人 × 4月、1日WORK / 2日LEAVE(REQUESTED) / 3〜30日REST。"""
    _, session = client
    d = Department(name="エクスポートテスト部門")
    session.add(d)
    await session.flush()
    session.add(WorkRuleConfig(department_id=d.id))

    emp = Employee(
        employee_number=44001, department_id=d.id, name="出力太郎",
        role="STAFF", available_days="0,1,2,3,4,5,6",
        available_shift_types="1,2,3",
    )
    session.add(emp)

    grp = WorkPatternGroup(department_id=d.id, name="A", is_auxiliary=False)
    session.add(grp)
    await session.flush()
    pat = WorkPattern(
        group_id=grp.id, pattern_name="A1", shift_type=1,
        shift_start="07:00", shift_end="16:00",
    )
    session.add(pat)
    await session.flush()

    sch = ShiftSchedule(
        department_id=d.id, year=2026, month=4,
        status="GENERATED", generation_attempt=1, is_active=True,
    )
    session.add(sch)
    await session.flush()

    session.add(ShiftAssignment(
        schedule_id=sch.id, employee_id=emp.id, date=date(2026, 4, 1),
        assignment_type="WORK", pattern_id=pat.id,
    ))
    session.add(ShiftAssignment(
        schedule_id=sch.id, employee_id=emp.id, date=date(2026, 4, 2),
        assignment_type="LEAVE", pattern_id=None,
    ))
    session.add(LeaveRequest(
        employee_id=emp.id, year=2026, month=4, day=2, leave_type="REQUESTED",
    ))
    for day in range(3, 31):
        session.add(ShiftAssignment(
            schedule_id=sch.id, employee_id=emp.id, date=date(2026, 4, day),
            assignment_type="REST", pattern_id=None,
        ))
    await session.flush()
    return sch


@pytest.mark.asyncio
async def test_export_xlsx_returns_bytes(client: tuple, schedule_full: ShiftSchedule):
    _, session = client
    data = await export_schedule_xlsx(session, schedule_full)
    assert isinstance(data, bytes)
    # xlsx は ZIP (PK\x03\x04)
    assert data[:4] == b"PK\x03\x04"


@pytest.mark.asyncio
async def test_export_xlsx_contents(client: tuple, schedule_full: ShiftSchedule):
    """xlsx を再パースして主要セルが正しいか確認。"""
    _, session = client
    import io
    from openpyxl import load_workbook

    data = await export_schedule_xlsx(session, schedule_full)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active

    # ヘッダ
    assert ws.cell(row=1, column=1).value == "氏名"
    assert ws.cell(row=1, column=2).value == "役職"
    # データ行
    assert ws.cell(row=2, column=1).value == "出力太郎"
    assert ws.cell(row=2, column=2).value == "STAFF"
    # 1日(列3) = WORK A1
    assert ws.cell(row=2, column=3).value == "A1"
    # 2日(列4) = LEAVE REQUESTED → ●
    assert ws.cell(row=2, column=4).value == "●"
    # 3日(列5) = REST → 割当休日
    assert ws.cell(row=2, column=5).value == "割当休日"
    # 勤務日数: 1, 休日数: 29 (4月30日中)
    assert ws.cell(row=2, column=33).value == 1
    assert ws.cell(row=2, column=34).value == 29


@pytest.mark.asyncio
async def test_export_csv_long(client: tuple, schedule_full: ShiftSchedule):
    _, session = client
    data = await export_schedule_csv_long(session, schedule_full)
    assert isinstance(data, bytes)
    text = data.decode("cp932")
    lines = [l for l in text.splitlines() if l]

    header = lines[0].split(",")
    assert header == ["社員番号", "氏名", "日付", "祝日", "作業パターン", "勤務区分"]
    # 1ヘッダ + 30日
    assert len(lines) == 31

    by_date = {row.split(",")[2]: row.split(",") for row in lines[1:]}
    # 1日: WORK A1
    assert by_date["2026/04/01"][4] == "A1"
    assert by_date["2026/04/01"][5] == "勤務"
    # 2日: LEAVE REQUESTED → コード 1
    assert by_date["2026/04/02"][4] == "1"
    assert by_date["2026/04/02"][5] == "休日"
    # 3日: REST → コード 4
    assert by_date["2026/04/03"][4] == "4"
    assert by_date["2026/04/03"][5] == "休日"


@pytest.mark.asyncio
async def test_export_csv_long_no_employees(client: tuple):
    """従業員ゼロの部門でも空行で正常に返る。"""
    _, session = client
    d = Department(name="空部門")
    session.add(d)
    await session.flush()
    session.add(WorkRuleConfig(department_id=d.id))
    sch = ShiftSchedule(
        department_id=d.id, year=2026, month=4,
        status="GENERATED", generation_attempt=1, is_active=True,
    )
    session.add(sch)
    await session.flush()

    data = await export_schedule_csv_long(session, sch)
    text = data.decode("cp932")
    lines = [l for l in text.splitlines() if l]
    # ヘッダのみ
    assert len(lines) == 1
